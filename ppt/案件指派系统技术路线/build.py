# -*- coding: utf-8 -*-
"""案件指派系统 · 技术路线汇报（简明版 12 页）

风格沿用 DESIGN.md 的红金政务风 + 公安警蓝主色。
画布 1280x720，母版：A 留白 0-48 / B 标题 48-110 / C 内容 110-660 / D 页脚 660-700。

版面硬规矩（上一版踩过的坑，这里全部遵守）：
  1. 文本框不会自动撑高：卡片高度一律按「行数 × 行高」手算，宁高勿低。
  2. 同一框内多段 para 会渲染到同一行：列表一律用**单个 para + \\n**（rich 支持）。
  3. 卡头文字用 MIDDLE（card 内部已处理），卡内正文用 TOP 对齐。
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ---------------- 设计令牌 ----------------
NAVY = RGBColor(0x12, 0x29, 0x4A)
BLUE = RGBColor(0x1B, 0x4A, 0x8C)
BLUE2 = RGBColor(0x2A, 0x5D, 0xA6)
BLUE_BG = RGBColor(0xE8, 0xEF, 0xF8)
GOLD = RGBColor(0xC8, 0xA4, 0x5C)
TEXT = RGBColor(0x1B, 0x24, 0x30)
TEXT2 = RGBColor(0x5A, 0x64, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC6, 0x2A, 0x2A)
GRAY = RGBColor(0x8A, 0x92, 0x9E)
LINE = RGBColor(0xDF, 0xE4, 0xEA)
ON_NAVY = RGBColor(0xB9, 0xCB, 0xE4)   # 藏蓝底上的次字

F_TITLE = '宋体'
F_BODY = '微软雅黑'
F_LATIN = 'Arial'

W, H = 1280, 720
PAD = 70
TOTAL = 12

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, 'assets')
OUT = os.path.join(BASE, '案件指派系统技术路线汇报.pptx')


# ---------------- 基础工具 ----------------
def px(v):
    return Emu(int(v * 9525))


def rect(slide, x, y, w, h, fill=None, line=None, lw=1,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size=22, color=TEXT, bold=False, font=F_BODY,
         align=PP_ALIGN.LEFT, space_before=0, space_after=6, line=1.35, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    parts = str(text).split('\n')
    for i, seg in enumerate(parts):
        if i:
            p.add_line_break()
        r = p.add_run()
        r.text = seg
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return p


def rich(tf, segs, size=22, align=PP_ALIGN.LEFT, space_before=0,
         space_after=6, line=1.35, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for text, st in segs:
        parts = str(text).split('\n')
        for i, seg in enumerate(parts):
            if i:
                p.add_line_break()
            r = p.add_run()
            r.text = seg
            r.font.size = Pt(st.get('size', size))
            r.font.color.rgb = st.get('color', TEXT)
            r.font.bold = st.get('bold', False)
            r.font.name = st.get('font', F_BODY)
    return p


def rows_segs(rows, size=19, key=BLUE, val=TEXT, val2=TEXT2):
    """rows = [(标签, 说明), ...] 或 [('', 单行文本), ...] → rich 的 segs"""
    segs = []
    for i, item in enumerate(rows):
        k, v = item if isinstance(item, tuple) else ('', item)
        pre = '' if i == 0 else '\n'
        if k:
            segs.append((pre + k + '　', {'size': size, 'bold': True, 'color': key}))
            segs.append((v, {'size': size, 'color': val}))
        else:
            segs.append((pre + v, {'size': size, 'color': val}))
    return segs


def put(slide, x, y, w, rows, size=19, lh=1.5, key=BLUE, val=TEXT):
    """单框写多行（一个 para + 软换行），TOP 对齐"""
    tf = tb(slide, x, y, w, 20)
    rich(tf, rows_segs(rows, size, key=key, val=val), size=size,
         line=lh, space_after=0, first=True)
    return tf


def put2(slide, x, y, w1, w2, left, right, size=18, lh=1.45, gap=18):
    put(slide, x, y, w1, left, size=size, lh=lh)
    put(slide, x + w1 + gap, y, w2, right, size=size, lh=lh)


def footer(slide, idx, dark=False):
    y = 664
    rect(slide, 0, y, W * 0.52, 9, fill=BLUE)
    rect(slide, W * 0.52, y, W * 0.30, 9, fill=NAVY)
    rect(slide, W * 0.82, y, W * 0.18, 9, fill=GOLD)
    c = RGBColor(0x9F, 0xB4, 0xD0) if dark else GRAY
    tf = tb(slide, PAD, y + 18, W - PAD * 2, 24)
    para(tf, '案件指派系统 CaseFlow · 技术路线汇报', size=14, color=c,
         space_after=0, first=True)
    tf2 = tb(slide, W - PAD - 160, y + 18, 160, 24)
    para(tf2, '%02d / %d' % (idx, TOTAL), size=14, color=c, align=PP_ALIGN.RIGHT,
         space_after=0, font=F_LATIN, first=True)


def title(slide, text, idx, sub=None):
    tf = tb(slide, PAD, 52, W - PAD * 2, 58)
    para(tf, text, size=34, color=NAVY, bold=True, font=F_TITLE,
         space_after=0, first=True)
    if sub:
        tf2 = tb(slide, PAD, 100, W - PAD * 2, 22)
        para(tf2, sub, size=16, color=TEXT2, space_after=0, first=True)
    footer(slide, idx)


def card(slide, x, y, w, h, head=None, kind='A', head_h=46):
    """A 白卡 / B 警蓝头白卡 / C 藏蓝实心卡"""
    if kind == 'C':
        rect(slide, x, y, w, h, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        return
    rect(slide, x, y, w, h, fill=WHITE, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    if head and kind == 'B':
        rect(slide, x, y, w, head_h, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        rect(slide, x, y + head_h - 10, w, 10, fill=BLUE)
        tf = tb(slide, x + 24, y, w - 48, head_h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, head, size=22, color=WHITE, bold=True, space_after=0, first=True)


def pic_fit(slide, name, x, y, box_w, box_h):
    """等比 contain 贴图，返回 (w, h)"""
    p = os.path.join(ASSETS, name)
    if not os.path.exists(p):
        return (0, 0)
    im = Image.open(p)
    ar = im.height / float(im.width)
    w, h = box_w, box_w * ar
    if h > box_h:
        h, w = box_h, box_h / ar
    slide.shapes.add_picture(p, px(x), px(y), px(w), px(h))
    return (w, h)


def hero_bg(slide):
    rect(slide, 0, 0, W, H, fill=NAVY)
    for r, a in ((300, RGBColor(0x1A, 0x36, 0x5E)), (220, RGBColor(0x22, 0x44, 0x74))):
        rect(slide, W - 300 - r / 2.0, 150 - r / 2.0, r, r, fill=a, shape=MSO_SHAPE.OVAL)
    for i in range(9):
        rect(slide, 40 + i * 46, 470, 2, 150, fill=RGBColor(0x1C, 0x3B, 0x66))


# ==================================================================
#  P01 封面
# ==================================================================
def p01(s):
    hero_bg(s)
    rect(s, PAD, 196, 6, 232, fill=GOLD)
    tf = tb(s, 104, 196, 900, 130)
    para(tf, '案件指派系统', size=92, color=WHITE, bold=True, font=F_TITLE,
         space_after=0, first=True)
    tf = tb(s, 108, 322, 700, 54)
    para(tf, 'CaseFlow', size=38, color=GOLD, bold=True, font=F_LATIN,
         space_after=0, first=True)
    rect(s, 104, 392, 486, 58, fill=BLUE)
    tf = tb(s, 128, 392, 440, 58, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, '技术路线汇报 · 简明版', size=25, color=WHITE, bold=True,
         space_after=0, first=True)
    for i, txt in enumerate(('汇报对象：技术评审', '2026 年 10 月')):
        x = 104 + i * 260
        rect(s, x, 478, 236, 42, fill=None, line=RGBColor(0x4A, 0x6E, 0x9E), lw=1.5,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28)
        tf = tb(s, x, 478, 236, 42, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, txt, size=17, color=RGBColor(0xD6, 0xE2, 0xF2), align=PP_ALIGN.CENTER,
             space_after=0, first=True)
    footer(s, 1, dark=True)


# ==================================================================
#  P02 目录
# ==================================================================
def p02(s):
    tf = tb(s, PAD, 56, 400, 74)
    para(tf, '目 录', size=56, color=NAVY, bold=True, font=F_TITLE,
         space_after=0, first=True)
    rect(s, PAD, 138, 64, 5, fill=GOLD)
    items = [
        ('01', '技术路线总览', '五层架构与请求链路', '03'),
        ('02', '后端技术栈', '框架与数据访问选型', '04 - 05'),
        ('03', '数据库设计', '17 张表 · 四大域', '06 - 07'),
        ('04', '接口与路由', '12 个控制器 · 两道拦截', '08 - 09'),
        ('05', '实时与权限', 'SSE 单一埋点 · 角色授权', '10'),
        ('06', '部署与运行', 'H2 开发 · MySQL 生产', '11'),
    ]
    cw, ch = 364, 176
    for i, (no, name, desc, page) in enumerate(items):
        col, row = i % 3, i // 3
        x = PAD + col * (cw + 24)
        y = 168 + row * (ch + 24)
        card(s, x, y, cw, ch, head=no + '　' + name, kind='B', head_h=46)
        put(s, x + 24, y + 64, cw - 48, [('', desc)], size=17, lh=1.5, val=TEXT2)
        rect(s, x + 24, y + 128, 40, 2, fill=GOLD)
        tf = tb(s, x + 24, y + 138, cw - 48, 24)
        para(tf, 'P ' + page, size=15, color=GRAY, space_after=0, font=F_LATIN, first=True)
    footer(s, 2)


# ==================================================================
#  P03 技术路线总览
# ==================================================================
def p03(s):
    title(s, '技术路线总览', 3, sub='一次请求从浏览器到数据库的完整链路，共五层')
    layers = [
        ('表现层', 'Vue3 + Element Plus + ECharts，纯桌面端内网浏览器'),
        ('接口层', 'REST + JSON，统一 /api 前缀，令牌校验'),
        ('业务层', 'Spring Boot 2.7 · 23 个 Service 承载全部业务规则'),
        ('数据访问层', 'MyBatis-Plus 3.5.5 · Mapper 与条件构造器'),
        ('数据层', 'H2 文件库（开发）/ MySQL 8（生产）'),
    ]
    lx, lw2 = PAD, 690
    y = 150
    for i, (nm, desc) in enumerate(layers):
        rect(s, lx, y, lw2, 66, fill=BLUE_BG if i % 2 == 0 else WHITE, line=LINE)
        rect(s, lx, y, 8, 66, fill=BLUE if i % 2 == 0 else BLUE2)
        tf = tb(s, lx + 24, y + 8, 200, 28)
        para(tf, nm, size=20, color=NAVY, bold=True, space_after=0, first=True)
        tf = tb(s, lx + 24, y + 36, lw2 - 48, 26)
        para(tf, desc, size=16, color=TEXT2, space_after=0, first=True)
        if i < len(layers) - 1:
            rect(s, lx + 24, y + 68, 44, 3, fill=GOLD)
        y += 78
    # 右侧：五道关卡（藏蓝实心卡）
    cx, cy, cwd = 790, 150, 420
    chh = y - 12 - cy
    card(s, cx, cy, cwd, chh, kind='C')
    tf = tb(s, cx + 26, cy + 22, cwd - 52, 36)
    para(tf, '一次请求经过的五道关卡', size=23, color=WHITE, bold=True,
         space_after=0, first=True)
    rect(s, cx + 26, cy + 66, 60, 3, fill=GOLD)
    gates = [
        ('① 路由守卫', '未登录跳登录页'),
        ('② 令牌校验', '解析核对登录状态'),
        ('③ 权限拦截', '管理层专属接口'),
        ('④ 业务校验', '服务层细粒度权限'),
        ('⑤ 落库广播', '写日志发通知推实时'),
    ]
    rows = [(a, b) for a, b in gates]
    put(s, cx + 26, cy + 86, cwd - 52, rows, size=17, lh=1.7,
        key=GOLD, val=RGBColor(0xE4, 0xEC, 0xF6))


# ==================================================================
#  P04 后端技术栈
# ==================================================================
def p04(s):
    title(s, '后端技术栈', 4,
          sub='选型原则：内网环境兼容性优先，稳定压过新特性 —— JDK 8 + Spring Boot 2.x')
    cw2, ch2 = 558, 220
    blocks = [
        ('运行底座', [
            ('JDK', '1.8（内网服务器环境限制）'),
            ('框架', 'Spring Boot 2.7.18'),
            ('容器', '内嵌 Tomcat，端口 8080'),
        ]),
        ('数据访问', [
            ('ORM', 'MyBatis-Plus 3.5.5'),
            ('写法', '条件构造器 + 少量手写 SQL'),
            ('分页', 'MyBatis-Plus 分页插件'),
        ]),
        ('通用能力', [
            ('口令', 'spring-security-crypto · BCrypt'),
            ('校验', 'spring-boot-starter-validation'),
            ('表格', 'EasyExcel 3.3.4 导入导出'),
        ]),
        ('工程规模', [
            ('源文件', '125 个 Java 文件'),
            ('代码量', '14,834 行'),
            ('结构', '12 个控制器 · 23 个服务'),
        ]),
    ]
    for i, (head, rows) in enumerate(blocks):
        col, row = i % 2, i // 2
        x = PAD + col * (cw2 + 24)
        y = 150 + row * (ch2 + 20)
        card(s, x, y, cw2, ch2, head=head, kind='B')
        put(s, x + 24, y + 64, cw2 - 48, rows, size=19, lh=1.75)
    tf = tb(s, PAD, 620, W - PAD * 2, 22)
    para(tf, '注：Spring Boot 2.x 与 JDK 8 为内网既有环境基线，升级到 3.x 需同步升级 JDK 17，暂不纳入。',
         size=15, color=GRAY, space_after=0, first=True)


# ==================================================================
#  P05 前端技术栈
# ==================================================================
def p05(s):
    title(s, '前端技术栈', 5, sub='Vue3 组合式 API + Element Plus，构建工具 Vite 5')
    cw2 = 558
    card(s, PAD, 150, cw2, 226, head='核心依赖', kind='B')
    put2(s, PAD + 24, 150 + 64, 244, 244,
         [('Vue', '3.4 组合式 API'),
          ('Element Plus', '2.6'),
          ('Pinia', '2.1 状态管理'),
          ('Vite', '5.2 构建热更新')],
         [('ECharts', '5.5.1 图表'),
          ('axios', '1.6 请求封装'),
          ('dayjs', '日期期限计算'),
          ('SortableJS', '拖拽排序')],
         size=17, lh=1.5, gap=24)
    card(s, 652, 150, cw2, 226, head='工程约定', kind='B')
    put(s, 652 + 24, 150 + 64, cw2 - 48, [
        ('规模', '33 个视图与组件 · 11,262 行'),
        ('状态', '跨组件计数与状态一律放 Pinia'),
        ('图表', '统一 EChart.vue 封装 + 统一主题'),
    ], size=19, lh=1.75)
    # 底部三条硬约定
    card(s, PAD, 404, W - PAD * 2, 234, head='三条硬约定（改代码前必读）', kind='B')
    put(s, PAD + 24, 404 + 64, W - PAD * 2 - 48, [
        ('图表', '新增图表类型必须先在 EChart.vue 注册，否则静默不渲染、零报错'),
        ('加字段', '统一走 SchemaMigration.NEW_COLUMNS，schema.sql 的建表语句不补列'),
        ('双副本', 'resources/sql 与 app/sql 两份 schema.sql 必须同步改'),
    ], size=19, lh=1.75)


# ==================================================================
#  P06 数据库设计
# ==================================================================
def p06(s):
    title(s, '数据库设计', 6,
          sub='17 张表分四大域；建表走 schema.sql，存量库补列走 SchemaMigration（MySQL 的 ALTER 不支持 IF EXISTS）')
    cw2, ch2 = 558, 230
    blocks = [
        ('域一　组织与账号', [
            ('sys_user', '账号、角色、审核状态'),
            ('org_employee', '员工档案、部门、组织层级'),
        ]),
        ('域二　案件主体', [
            ('case_info / case_assignee', '案件本体与承办协办人'),
            ('case_file / case_suspect', '附件与涉案人员'),
            ('case_category / case_plan', '类别与流程任务'),
            ('case_approval', '审批记录'),
        ]),
        ('域三　待办与意见', [
            ('case_todo / case_todo_feedback', '两级待办与累积反馈'),
            ('case_leader_opinion / case_question', '领导意见与疑问问答'),
            ('case_opinion_read', '意见已读（登录人维度）'),
        ]),
        ('域四　追溯与通知', [
            ('operation_log', '操作日志与双快照'),
            ('case_notification', '统一信箱'),
            ('case_progress_comment', '进展说明'),
        ]),
    ]
    for i, (head, rows) in enumerate(blocks):
        col, row = i % 2, i // 2
        x = PAD + col * (cw2 + 24)
        y = 150 + row * (ch2 + 20)
        card(s, x, y, cw2, ch2, head=head, kind='B')
        put(s, x + 24, y + 64, cw2 - 48, rows, size=17, lh=1.6)


# ==================================================================
#  P07 三张核心表
# ==================================================================
def p07(s):
    title(s, '三张核心表', 7, sub='看懂这三张表，就看懂了系统的主干')
    cw3 = 364
    blocks = [
        ('案件主体', 'case_info', [
            ('flow_stage', '流程阶段四态'),
            ('status', '状态与办结标记'),
            ('assignee', '主办协办分离'),
            ('期限', '驱动到期提醒'),
        ]),
        ('待办', 'case_todo', [
            ('parent_id', '两级主子任务'),
            ('opinion_id', '由意见派生'),
            ('status', '落实状态校验'),
            ('due_date', '待办级期限'),
        ]),
        ('操作日志', 'operation_log', [
            ('snapshot_before', '写前快照'),
            ('snapshot_after', '写后快照'),
            ('undone', '已撤回标记'),
            ('undo_of', '指向原日志'),
        ]),
    ]
    for i, (head, table, rows) in enumerate(blocks):
        x = PAD + i * (cw3 + 24)
        card(s, x, 150, cw3, 300, head=head, kind='B', head_h=46)
        tf = tb(s, x + 22, 150 + 60, cw3 - 44, 24)
        para(tf, table, size=16, color=GRAY, bold=True, font=F_LATIN,
             space_after=0, first=True)
        put(s, x + 22, 150 + 92, cw3 - 44, rows, size=16, lh=1.55)
    card(s, PAD, 478, W - PAD * 2, 160, kind='C')
    tf = tb(s, PAD + 26, 500, W - PAD * 2 - 52, 36)
    para(tf, '撤回机制：一次写操作前后各存一份完整快照', size=23, color=WHITE,
         bold=True, space_after=0, first=True)
    rect(s, PAD + 26, 542, 60, 3, fill=GOLD)
    tf = tb(s, PAD + 26, 556, W - PAD * 2 - 52, 62)
    para(tf, '撤回 = 把写前快照原样写回；撤回动作本身也记一条日志，再撤一次即重做。'
             '仅案件类可撤，且必须是该案件最新一条。',
         size=17, color=ON_NAVY, space_after=0, line=1.45, first=True)


# ==================================================================
#  P08 接口与路由（后端）
# ==================================================================
def p08(s):
    title(s, '接口与路由', 8, sub='统一 /api 前缀 · 12 个控制器 · 两道拦截')
    card(s, PAD, 150, 558, 298, head='控制器清单', kind='B')
    put2(s, PAD + 24, 214, 236, 236,
         [('Auth', '登录注册'),
          ('Case', '案件'),
          ('Category', '类别'),
          ('Employee', '员工'),
          ('Event', '实时流'),
          ('File', '附件')],
         [('Log', '操作日志'),
          ('Notification', '通知'),
          ('Question', '疑问'),
          ('Todo', '待办'),
          ('User', '账号'),
          ('Watch', '盯办')],
         size=17, lh=1.45, gap=32)
    card(s, PAD, 460, 558, 178, head='统一约定', kind='B')
    put(s, PAD + 24, 524, 510, [
        ('前缀', '统一 /api，网关与前端同源'),
        ('响应', '统一 JSON 结构与异常处理'),
        ('规模', '132 处接口映射'),
    ], size=17, lh=1.5)
    card(s, 652, 150, 558, 240, head='两道拦截', kind='B')
    put(s, 652 + 24, 214, 510, [
        ('第一道', 'LoginInterceptor：请求头优先、query 令牌兜底'),
        ('第二道', 'PermissionInterceptor：注解方法管理层专属'),
        ('覆盖', '@FullAccessOnly 共 32 处'),
    ], size=18, lh=1.6)
    card(s, 652, 410, 558, 228, head='分层授权惯例', kind='B')
    put(s, 652 + 24, 474, 510, [
        ('主任务', '领导定的清单，仅管理层可改'),
        ('子任务', '承办人自己拆的，承办人即可'),
        ('要点', '注解在方法执行前拦截，写在服务层无效'),
    ], size=18, lh=1.6)


# ==================================================================
#  P09 前端页面与路由守卫
# ==================================================================
def p09(s):
    title(s, '前端页面与路由守卫', 9, sub='13 个视图 · 6 个状态仓库 · 三条守卫规则')
    card(s, PAD, 150, 558, 334, head='页面清单', kind='B')
    put2(s, PAD + 24, 214, 236, 236,
         [('', '我的案件'), ('', '我的待办'), ('', '工作台'),
          ('', '案件管理'), ('', '到期提醒'), ('', '员工图谱'), ('', '案件盯办')],
         [('', '待办总览'), ('', '账号管理'), ('', '类别管理'),
          ('', '选择案件类型'), ('', '登录'), ('', '注册')],
         size=17, lh=1.45, gap=32)
    card(s, PAD, 504, 558, 134, head='状态仓库', kind='B')
    put(s, PAD + 24, 568, 510, [
        ('', 'user · events · caseType · category · myTodo · pending'),
    ], size=17, lh=1.5, val=TEXT2)
    card(s, 652, 150, 558, 220, head='路由守卫三条', kind='B')
    put(s, 652 + 24, 214, 510, [
        ('登录', '未登录一律跳登录页'),
        ('权限', '管理层页面对非管理层拦截'),
        ('门控', '未选案件类型跳类型选择页'),
    ], size=18, lh=1.6)
    w, h = pic_fit(s, 'ui-dashboard.png', 652, 392, 558, 216)
    rect(s, 652, 392, w, h, fill=None, line=LINE)
    tf = tb(s, 652, 620, 558, 22)
    para(tf, '系统工作台实拍：待办、期限与统计一屏可见', size=15, color=GRAY,
         space_after=0, first=True)


# ==================================================================
#  P10 实时推送与权限模型
# ==================================================================
def p10(s):
    title(s, '实时推送与权限模型', 10, sub='全站埋点只有一个汇聚点：写日志')
    seg_w = (W - PAD * 2 - 4 * 18) / 5.0
    x = PAD
    chain = [('业务写操作', '服务层落库'), ('统一埋点', 'LogService.log'),
             ('广播', 'SseHub 推送'), ('前端订阅', 'store/events.js'),
             ('页面刷新', '局部重取数据')]
    for i, (a, b) in enumerate(chain):
        rect(s, x, 150, seg_w, 100, fill=WHITE, line=LINE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        rect(s, x, 150, seg_w, 6, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tf = tb(s, x + 16, 168, seg_w - 32, 28)
        para(tf, a, size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
             space_after=0, first=True)
        tf = tb(s, x + 16, 198, seg_w - 32, 30)
        para(tf, b, size=15, color=TEXT2, align=PP_ALIGN.CENTER, space_after=0, first=True)
        if i < 4:
            rect(s, x + seg_w + 2, 192, 14, 14, fill=GOLD, shape=MSO_SHAPE.RIGHT_ARROW)
        x += seg_w + 18
    card(s, PAD, 282, 558, 200, head='角色体系', kind='B')
    put(s, PAD + 24, 346, 510, [
        ('全权限', '局长 · 大队长 · 副大队长 · 法制员'),
        ('受限', '民警：仅本人相关与公开栏目'),
    ], size=19, lh=1.7)
    card(s, 652, 282, 558, 200, head='通知分发口径', kind='B')
    put(s, 652 + 24, 346, 510, [
        ('普通用户', '本案在办的承办人与协办人'),
        ('管理层', '全站所有用户（除自己）'),
        ('跳转', '带锚点精确定位并高亮目标'),
    ], size=19, lh=1.7)
    card(s, PAD, 502, W - PAD * 2, 136, kind='C')
    tf = tb(s, PAD + 26, 522, W - PAD * 2 - 52, 34)
    para(tf, '为什么埋点只有一处', size=22, color=WHITE, bold=True,
         space_after=0, first=True)
    rect(s, PAD + 26, 562, 60, 3, fill=GOLD)
    tf = tb(s, PAD + 26, 574, W - PAD * 2 - 52, 52)
    para(tf, '所有写操作统一走写日志方法，落库后自动触发通知与实时广播，'
             '新功能即自动获得三件套能力。',
         size=17, color=ON_NAVY, line=1.45, space_after=0, first=True)


# ==================================================================
#  P11 部署与运行
# ==================================================================
def p11(s):
    title(s, '部署与运行', 11, sub='开发态一键启动，生产态打成一体包走内网')
    cw3 = 364
    blocks = [
        ('开发环境', [
            ('数据库', 'H2 文件库，零安装'),
            ('启动', 'dev.py 一键拉起'),
            ('端口', '后端 8080 / 前端 5173'),
            ('自检', '8 个脚本覆盖核心流程'),
        ]),
        ('生产部署', [
            ('数据库', '本地 MySQL 8 实例'),
            ('打包', 'dist 打进后端 static'),
            ('访问', '内网 IP:8080/api'),
            ('脚本', '防火墙放行 + 二维码'),
        ]),
        ('注意事项', [
            ('H2', '同时只能一个后端进程'),
            ('防火墙', '放行 8080 入站'),
            ('脚本', '批处理须 GBK + CRLF'),
            ('库文件', '演示库不进版本库'),
        ]),
    ]
    for i, (head, rows) in enumerate(blocks):
        x = PAD + i * (cw3 + 24)
        card(s, x, 150, cw3, 330, head=head, kind='B')
        put(s, x + 22, 214, cw3 - 44, rows, size=18, lh=1.6)
    card(s, PAD, 500, W - PAD * 2, 138, kind='C')
    tf = tb(s, PAD + 26, 520, W - PAD * 2 - 52, 34)
    para(tf, '一条命令启动', size=22, color=WHITE, bold=True, space_after=0, first=True)
    rect(s, PAD + 26, 560, 60, 3, fill=GOLD)
    tf = tb(s, PAD + 26, 572, W - PAD * 2 - 52, 52)
    para(tf, '.venv\\Scripts\\python.exe scripts\\dev.py --seed　'
             '（加 --profile mysql 切生产库，--check 先做环境自检）',
         size=17, color=ON_NAVY, font=F_LATIN, line=1.45, space_after=0, first=True)


# ==================================================================
#  P12 结束页
# ==================================================================
def p12(s):
    hero_bg(s)
    rect(s, PAD, 236, 6, 190, fill=GOLD)
    tf = tb(s, 104, 236, 900, 130)
    para(tf, '汇报完毕', size=76, color=WHITE, bold=True, font=F_TITLE,
         space_after=0, first=True)
    tf = tb(s, 108, 372, 700, 40)
    para(tf, '案件指派系统 · 技术路线', size=26, color=GOLD, bold=True,
         space_after=0, first=True)
    nums = [('17', '张数据表'), ('12', '个控制器'), ('2.6', '万行代码')]
    x = 104
    for v, lab in nums:
        tf = tb(s, x, 440, 220, 72)
        para(tf, v, size=48, color=WHITE, bold=True, font=F_LATIN,
             space_after=0, line=1.1, first=True)
        tf = tb(s, x, 520, 220, 30)
        para(tf, lab, size=18, color=ON_NAVY, space_after=0, first=True)
        x += 250
    footer(s, 12, dark=True)


# ==================================================================
def main():
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    blank = prs.slide_layouts[6]
    for fn in (p01, p02, p03, p04, p05, p06, p07, p08, p09, p10, p11, p12):
        fn(prs.slides.add_slide(blank))
    prs.save(OUT)
    print('已生成：%s（%d 页）' % (OUT, len(prs.slides.__iter__.__self__._sldIdLst)))


if __name__ == '__main__':
    main()
