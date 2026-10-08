# -*- coding: utf-8 -*-
"""版面体检：直接读 .pptx，用 PIL 实测字体度量算真实渲染高度，
比对「内容区下沿 656」与「所属卡片底边」，报出血与超卡。"""
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

BASE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(BASE, '案件指派系统技术路线汇报.pptx')
EMU_PX = 9525.0
C_BOTTOM = 656.0

FONTS = {
    ('宋体', True): ('C:/Windows/Fonts/simhei.ttf', 0),
    ('宋体', False): ('C:/Windows/Fonts/simsun.ttc', 0),
    ('微软雅黑', True): ('C:/Windows/Fonts/msyhbd.ttc', 0),
    ('微软雅黑', False): ('C:/Windows/Fonts/msyh.ttc', 0),
    ('Arial', True): ('C:/Windows/Fonts/arialbd.ttf', 0),
    ('Arial', False): ('C:/Windows/Fonts/arial.ttf', 0),
}
_cache = {}


def font(name, bold, size_pt):
    pxv = max(6, int(round(size_pt * 96.0 / 72.0)))
    key = (name, bold, pxv)
    if key in _cache:
        return _cache[key]
    path, idx = FONTS.get((name, bold)) or FONTS[('微软雅黑', bold)]
    f = ImageFont.truetype(path, pxv, index=idx)
    _cache[key] = f
    return f


def e2p(v):
    return (v or 0) / EMU_PX


def run_width_px(text, name, bold, size_pt):
    if not text:
        return 0.0
    return font(name, bold, size_pt).getlength(text)


def frame_metrics(tf, box_w):
    total = 0.0
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            total += 13 * 96.0 / 72.0 * 1.2
            continue
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
        sa = p.space_after.pt if p.space_after else 0
        sb = p.space_before.pt if p.space_before else 0
        size_pt = max((r.font.size.pt for r in runs if r.font.size), default=13.0)
        cur = 0.0
        lines = 1
        for r in runs:
            segs = r.text.split('\n')
            for i, seg in enumerate(segs):
                if i:
                    total += lines * size_pt * 96.0 / 72.0 * ls
                    lines = 1
                    cur = 0.0
                w = run_width_px(seg, r.font.name or '微软雅黑',
                                 bool(r.font.bold), size_pt)
                cur += w
                if cur > box_w:
                    lines += int(cur // box_w)
                    cur = cur % box_w
        total += lines * size_pt * 96.0 / 72.0 * ls
        total += (sa + sb) * 96.0 / 72.0
    return total


prs = Presentation(PPTX)
print('画布 %.0f × %.0f，共 %d 页' % (e2p(prs.slide_width), e2p(prs.slide_height),
                                     len(prs.slides._sldIdLst)))
print('=' * 88)
issues = []
for si, slide in enumerate(prs.slides, 1):
    cards = []
    texts = []
    for sh in slide.shapes:
        x, y = e2p(sh.left), e2p(sh.top)
        w, h = e2p(sh.width), e2p(sh.height)
        if sh.has_text_frame:
            t = ''.join(r.text for p in sh.text_frame.paragraphs for r in p.runs).strip()
            if t:
                texts.append((x, y, w, h, sh.text_frame, t))
        elif w > 120 and h > 60:
            cards.append((x, y, w, h))
    for (x, y, w, h, tf, t) in texts:
        if '技术路线汇报' in t or '/ 12' in t:
            continue
        need = frame_metrics(tf, w)
        bottom = y + need
        if bottom > C_BOTTOM + 4 and y < C_BOTTOM:
            issues.append((si, '出血', t[:30], y, bottom, C_BOTTOM))
        for (cx, cy, cw, ch) in cards:
            if cx - 4 <= x <= cx + cw and cy - 4 <= y <= cy + ch and bottom > cy + ch + 3:
                issues.append((si, '超卡', t[:30], bottom, cy + ch, (cy + ch) - y))
                break

if not issues:
    print('未发现文字出血或超出卡片')
else:
    print('发现 %d 处：\n' % len(issues))
    cur = None
    for it in issues:
        if it[0] != cur:
            print('-- 第 %02d 页 --' % it[0])
            cur = it[0]
        if it[1] == '出血':
            print('   [出血] 「%s」 底边 %.0f 超出内容区 %.0f（超出 %.0fpx）'
                  % (it[2], it[4], it[5], it[4] - it[5]))
        else:
            print('   [超卡] 「%s」 底边 %.0f 超出卡底 %.0f（超出 %.0fpx）'
                  % (it[2], it[3], it[4], it[3] - it[4]))
    sys.exit(1)
