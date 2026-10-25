# -*- coding: utf-8 -*-
"""生成《CaseFlow 功能更新说明（2026-10-18 ~ 10-25）》Word 文档。

图片由 scripts/cf-shots-updates.mjs 抓取，输出到 docs/manual-v4/shots/。
用法（用系统 Python 3.10，它带 python-docx）：
  C:\\Users\\adamin\\AppData\\Local\\Programs\\Python\\Python310\\python.exe scripts\\gen_updates_doc.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # 仓库根（app/scripts → 上两级）
SHOTS = os.path.join(ROOT, 'docs', 'manual-v4', 'shots')
OUT = os.path.join(ROOT, 'docs', 'manual-v4', 'CaseFlow功能更新说明_2026-10-18至10-25.docx')

NAVY = RGBColor(0x12, 0x29, 0x4A)
BLUE = RGBColor(0x1B, 0x4A, 0x8C)
GOLD = RGBColor(0xC8, 0xA4, 0x5C)
GREY = RGBColor(0x5A, 0x64, 0x72)
RED = RGBColor(0xC0, 0x39, 0x2B)


def cn(run, font='宋体', size=10.5, bold=False, color=None):
    """设置中英文字体（中文必须写 eastAsia，否则 Word 里回落成宋体/乱码级差异）"""
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if color is not None:
        run.font.color.rgb = color
    return run


def para(doc, text='', size=10.5, bold=False, color=None, align=None,
         space_before=0, space_after=6, line=1.5, indent=0, font='宋体'):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(indent)
    if align is not None:
        p.alignment = align
    if text:
        cn(p.add_run(text), font=font, size=size, bold=bold, color=color)
    return p


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), hex_color)
    tcPr.append(sh)


def h1(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    if num:
        cn(p.add_run(num + '  '), font='微软雅黑', size=15, bold=True, color=GOLD)
    cn(p.add_run(text), font='微软雅黑', size=15, bold=True, color=NAVY)
    # 标题下压一条深蓝细线
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '3')
    bottom.set(qn('w:color'), '1B4A8C')
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    cn(p.add_run('▍'), font='微软雅黑', size=12, bold=True, color=BLUE)
    cn(p.add_run(text), font='微软雅黑', size=12, bold=True, color=NAVY)
    return p


def figure(doc, name, caption, width_cm=15.0):
    path = os.path.join(SHOTS, name)
    if not os.path.exists(path):
        para(doc, f'[缺图：{name}]', color=RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cn(cap.add_run(caption), font='微软雅黑', size=8.5, color=GREY)


def steps(doc, items):
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.8)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.5
        cn(p.add_run(f'{i}. '), font='微软雅黑', size=10.5, bold=True, color=BLUE)
        cn(p.add_run(t), size=10.5)


def bullets(doc, items, color=None):
    for t in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.8)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.5
        cn(p.add_run('· '), size=10.5, bold=True, color=color or BLUE)
        cn(p.add_run(t), size=10.5)


def tip(doc, text, kind='tip'):
    """提示条：浅底单元格，比正文更醒目"""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(15.8)
    shade(cell, 'F4F7FB' if kind == 'tip' else 'FDF3F2')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.4
    cn(p.add_run('提示  ' if kind == 'tip' else '注意  '),
       font='微软雅黑', size=9.5, bold=True, color=BLUE if kind == 'tip' else RED)
    cn(p.add_run(text), size=9.5, color=GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()

# ---------- 页面与默认样式 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.4)
sec.left_margin = sec.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# ---------- 封面 ----------
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cn(p.add_run('案件指派系统  CaseFlow'), font='微软雅黑', size=13, color=GREY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
cn(p.add_run('功能更新说明'), font='微软雅黑', size=32, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
cn(p.add_run('2026 年 10 月 18 日 — 10 月 25 日'), font='微软雅黑', size=13, color=GOLD)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cn(p.add_run('本次共 9 项功能改动，其中新增 6 项、调整优化 3 项'), font='微软雅黑', size=10.5, color=GREY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cn(p.add_run('数据结构只增不改，升级不影响已录入的案件与人口数据'), font='微软雅黑', size=10.5, color=GREY)

for _ in range(8):
    doc.add_paragraph()

tb = doc.add_table(rows=4, cols=2)
tb.alignment = WD_TABLE_ALIGNMENT.CENTER
tb.style = 'Table Grid'
tb.autofit = False
info = [('文档用途', '向使用单位说明本次升级新增/修改了什么、怎么用'),
        ('适用版本', '升级补丁 20261025 及之后'),
        ('编写日期', '2026 年 10 月 25 日'),
        ('配套材料', '《升级说明.txt》《部署教程》《用户使用手册》')]
for i, (k, v) in enumerate(info):
    c0, c1 = tb.rows[i].cells
    c0.width, c1.width = Cm(3.2), Cm(12.0)
    shade(c0, 'EEF2F8')
    cn(c0.paragraphs[0].add_run(k), font='微软雅黑', size=9.5, bold=True, color=NAVY)
    cn(c1.paragraphs[0].add_run(v), size=9.5)

doc.add_page_break()

# ---------- 目录 ----------
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
cn(p.add_run('本 文 内 容'), font='微软雅黑', size=14, bold=True, color=NAVY)

toc = [
    ('一', '本次升级总览', '9 项改动一览表：改了什么、谁会用'),
    ('二', '功能详解与使用方法', '逐项说明操作步骤，附界面截图'),
    ('', '　1. 强制措施变更：申请—审批', '承办人提交、领导批准后才生效'),
    ('', '　2. 强制措施到期提醒', '工作台与我的案件双入口'),
    ('', '　3. 嫌疑人信息可编辑', '身份证、电话录错可改'),
    ('', '　4. 案件名称与信息可编辑', '列表与详情两个入口'),
    ('', '　5. 列表改显示报案/受案时间', '可按时间排序'),
    ('', '　6. 领导确认可退回', '退回须写意见，自动派待办'),
    ('', '　7. 所有时间精确到天', '不再有时分秒'),
    ('', '　8. 待办卡片更紧凑', '一屏多看 2—3 条'),
    ('三', '升级须知：不动已录入数据', '只增不改、一次性闸门、只换程序'),
    ('四', '常见问题', '6 个高频疑问'),
]
for num, title, desc in toc:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    if num:
        cn(p.add_run(num + '、'), font='微软雅黑', size=11, bold=True, color=GOLD)
        cn(p.add_run(title), font='微软雅黑', size=11, bold=True, color=NAVY)
    else:
        cn(p.add_run(title), font='微软雅黑', size=10, color=NAVY)
    cn(p.add_run('    ' + desc), size=9, color=GREY)

doc.add_page_break()

# ---------- 一、总览 ----------
h1(doc, '本次升级总览', '一')

para(doc, '这 8 天的改动可以归成一句话：让“措施到期”这件事有人盯、有流程、有记录，'
          '同时把录入时容易出错的地方补上（案件名能改、嫌疑人能改、时间不再带时分秒）。',
     space_before=4)

rows = [
    ('1', '强制措施变更需领导审批', '新增', '承办人 · 领导', '换措施不再直接生效，领导批了才算'),
    ('2', '强制措施到期提醒', '新增', '领导 · 承办人', '到期前 7 天就在首页提醒，可直接登记下一措施'),
    ('3', '嫌疑人信息可编辑', '新增', '承办人', '录错了改就行，不用删了重录'),
    ('4', '案件可编辑、可改名', '新增', '管理层', '名称、案情填错可随时更正'),
    ('5', '领导确认可退回', '新增', '领导', '不认可完成情况时退回并写明原因'),
    ('6', '登记措施按角色分流', '新增', '领导', '领导自己登记直接生效，承办人登记走审批'),
    ('7', '列表改显示报案/受案时间', '调整', '所有人', '一眼看到案子什么时候进来，可按时间排序'),
    ('8', '时间精度统一到“天”', '调整', '所有人', '所有日期只到年月日，不再出现时分秒'),
    ('9', '待办卡片紧凑化', '优化', '所有人', '一屏能放下更多待办，不用来回滚'),
]
t = doc.add_table(rows=1, cols=5)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
heads = ['#', '功能', '类型', '谁会用', '一句话说明']
widths = [Cm(0.8), Cm(4.9), Cm(1.4), Cm(2.8), Cm(5.9)]
for i, h in enumerate(heads):
    cell = t.rows[0].cells[i]
    cell.width = widths[i]
    shade(cell, '12294A')
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cn(p.add_run(h), font='微软雅黑', size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
for r in rows:
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].width = widths[i]
        p = cells[i].paragraphs[0]
        if i in (0, 2, 3):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cn(p.add_run(v), size=9.5,
           bold=(i == 1), color=NAVY if i == 1 else None,
           font='微软雅黑' if i == 1 else '宋体')
    shade(cells[2], 'EAF1FA' if r[2] == '新增' else 'F6F3EA')

doc.add_page_break()

# ---------- 二、功能详解 ----------
h1(doc, '功能详解与使用方法', '二')

# 2.1
h2(doc, '1. 强制措施变更：从“自己改”变成“申请—审批”')
para(doc, '过去承办人点一下就把措施改了，领导事后才知道；措施直接关系办案期限，'
          '改错一天就可能超期。现在分两种人走两条路：', indent=21)
bullets(doc, [
    '领导（所长、副所长、法制员、系统管理员）登记 → 立即生效，与以前一样快；',
    '承办人（办案民警）登记 → 只生成一份“待审批申请”，措施暂不改变，'
    '必须等领导同意后才真正生效。',
])
para(doc, '操作步骤：', bold=True, space_before=6)
steps(doc, [
    '在“工作台”或“我的案件”点“措施到期”卡片，右侧清单里找到那个案子；',
    '点“登记下一措施”，弹窗里选择措施类型（拘留 / 逮捕 / 取保候审 / 监视居住 / 拘传）；',
    '填“措施日期”，需要多长填“期限届满”（留空则按默认推算：拘留 3 天、逮捕 2 个月、取保 1 年、监居 6 个月）；',
    '备注可写变更理由，提交审批时会一起带给领导；',
    '点“确定”。承办人提交后按钮变成“审批中”，措施本身不变；领导在“待审批”里点“同意”后才落地，'
    '点“退回”则必须写意见。',
])
tip(doc, '弹窗底部那句“保存后生成待审批申请，措施暂不变，需领导批准后才生效”就是这套规则的说明，'
         '看到它就说明当前账号是承办人身份。')
figure(doc, '07-measure-dialog.png', '图 2-1  登记 / 变更强制措施弹窗（承办人提交后进入审批）', 10.5)

# 2.2
h2(doc, '2. 强制措施到期提醒：到期前 7 天就摆到桌面上')
para(doc, '措施一到期，羁押就失去合法依据，必须马上登记下一个措施。过去靠人记，现在系统自己算：', indent=21)
bullets(doc, [
    '工作台新增“措施到期”卡片，数字 = 已逾期 + 7 天内到期；',
    '点开卡片右侧列出清单：案件名称、当前措施、起算日、届满日、状态；',
    '状态为“已到期 N 天”的行，直接给“登记下一措施”按钮，不用再进案件；',
    '“我的案件”页也有同样的卡片和清单——民警看不到工作台，但名下案子的措施到期必须能看到。',
])
figure(doc, '01b-kpi-measure.png', '图 2-2  工作台顶部的“措施到期”卡片（红色数字 = 已逾期 + 7 天内到期的总数）', 7.5)
figure(doc, '03b-mine-closeup.png', '图 2-3  清单每行：案件名称 + 当前措施 + 届满日 + 到期状态；已到期的行右侧还有“登记下一措施”', 15.5)
figure(doc, '02-dashboard-measure-alert.png', '图 2-4  工作台上的同一份清单（“我的案件”页面按人过滤，只看自己名下的案子）', 15.5)

# 2.3
h2(doc, '3. 嫌疑人信息可编辑：录错不用删了重录')
para(doc, '以前嫌疑人一旦录进去，身份证填错、电话换了，只能删掉整条重录，录入人和录入时间也跟着丢。'
          '现在每一行右侧都有“编辑”。', indent=21)
para(doc, '操作步骤：', bold=True, space_before=6)
steps(doc, [
    '打开案件详情，滚到“嫌疑人”面板；',
    '点该行右侧的“编辑”（表格可左右滑动，按钮在最右边）；',
    '弹窗里会带回全部原值：姓名、性别、身份证号、联系电话、住址、强制措施及起止日、备注；',
    '改完点“保存”，列表立刻刷新，修改记录进操作日志，可撤回。',
])
tip(doc, '把“强制措施”清空成“未采取”再保存，起止时间会自动一并清掉——不会再出现“措施没了但日期还挂着”的矛盾数据。')
figure(doc, '05-drawer-suspect.png', '图 2-5  案件详情里的嫌疑人面板，措施与起止时间同列显示', 15.5)
figure(doc, '05c-suspect-actions.png', '图 2-6  表格滑到最右侧的“编辑 / 删除”按钮（窄屏时需横向滑动）', 7.0)
figure(doc, '06-suspect-dialog.png', '图 2-7  编辑嫌疑人弹窗：原值全部回显，措施与起止日各占整行不易填错', 10.5)

# 2.4
h2(doc, '4. 案件名称与信息可编辑')
para(doc, '新增、导入案件时名称打错、案情写得不准，以前只能作废重建。现在有两个入口可改：', indent=21)
bullets(doc, [
    '案件列表：每行操作列的“编辑”；',
    '案件详情：右上角“编辑案件”。',
])
para(doc, '打开后与建案表单一致，改完保存即生效，已录入的嫌疑人、意见、待办全部保留。', space_before=2)

# 2.5
h2(doc, '5. 列表改为显示“报案 / 受案时间”')
para(doc, '原来列表第一列是一串看不出含义的编号。现在换成时间：导入的案件按“受案时间”、'
          '手工新建的按“报案时间”，并在日期上方用小标签区分两者。', indent=21)
bullets(doc, [
    '这一列可以点表头排序，默认由新到旧，想找老案子切一下就行；',
    '案件编号仍然保留在系统里（导出、搜索、交接时用），只是不再占第一列。',
])
figure(doc, '04b-list-closeup.png', '图 2-8  案件列表首列改为“报案 / 受案时间”，标签区分导入件与手工建案', 15.5)

# 2.6
h2(doc, '6. 领导确认可“退回”')
para(doc, '办结需要领导确认。过去只有“同意”，不满意只能口头让人改。现在确认面板多了“退回”：', indent=21)
bullets(doc, [
    '退回必须写意见，意见自动落成一条待办派给承办人；',
    '案件阶段自动回到上一阶段，承办人在“我的待办”里就能看到要补什么；',
    '退回记录留在意见列表，可追溯是谁、什么时候、因为什么退回的。',
])

# 2.7
h2(doc, '7. 所有时间精确到“天”')
para(doc, '办案时间按天算，填到时分秒全是干扰。本次统一：', indent=21)
bullets(doc, [
    '所有日期控件只选年月日，列表、详情、提醒一律显示 2026-10-25 这种格式；',
    '“今天到期”按当天 23:59:59 判定，不会上午就被标成逾期；',
    '唯一例外是操作日志——审计要求，仍精确到秒。',
])

# 2.8
h2(doc, '8. 待办卡片更紧凑')
para(doc, '原来一张待办卡片要占大半屏，一页看不到几条。现在操作按钮横向排列、'
          '意见设置并进信息行，同样的屏幕能多放 2—3 条。', indent=21)
figure(doc, '05b-drawer-full.png', '图 2-9  案件详情全景：右下“案件待办”卡片已收紧，操作按钮横排', 15.5)

doc.add_page_break()

# ---------- 三、升级须知 ----------
h1(doc, '升级须知：不会动你已经录入的数据', '三')

para(doc, '使用单位最关心的问题：升级会不会把内网里的案子弄乱？答案是不会，'
          '本次在代码层面做了三条硬约束。', space_before=4)

h2(doc, '1. 数据结构只增不改')
para(doc, '升级只做两件事：缺表就建（表已存在则跳过）、缺列就补（新列一律允许为空）。', indent=21)
bullets(doc, [
    '不删表、不删列、不改字段类型、不改表名；',
    '本次新增的列全部可空，老数据行不需要任何转换就能通过；',
    '已用脚本核查：建表脚本与启动迁移中不存在 DROP / RENAME / MODIFY / TRUNCATE 语句。',
])

h2(doc, '2. 历史数据整理只跑一次')
para(doc, '以前每次启动都会做一些“数据整理”，反复执行会把人工调整过的内容改回去。'
          '现在这类操作全部加了一次性闸门：执行一次成功后写台账，之后永远跳过。', indent=21)
bullets(doc, [
    '台账表 schema_migration_mark 只记录“哪一步做过了”，与业务表无关；',
    '台账写不进去就不执行——宁可少整理，也不反复改数据；',
    '实测：第二次启动时一次性整理执行 0 次。',
])

h2(doc, '3. 升级只换程序文件')
para(doc, '发的是“升级补丁包”，里面只有程序（jar）和升级说明，不含配置目录与数据库目录，'
          '从物理上避免“整个解压覆盖”把口令改回占位、把数据目录清空。', indent=21)
steps(doc, [
    '停掉正在运行的服务；',
    '把新 jar 覆盖到原位置（配置文件、data 目录一律不动）；',
    '重新启动，程序自动补齐新增字段；',
    '打开页面确认能登录、能看到历史案件即可。',
])
tip(doc, '升级前后建议各备份一次数据库目录。实测同一套数据升级前后逐表比对：案件、嫌疑人、'
         '附件、待办、意见、审批全部原样，只有字典表按预期新增了行。')

doc.add_page_break()

# ---------- 四、常见问题 ----------
h1(doc, '常见问题', '四')

qa = [
    ('承办人登记了下一措施，为什么措施没变？',
     '这是本次的审批规则：承办人提交后只是“待审批申请”，措施要等领导同意才生效。'
     '在待审批期间，按钮会显示“审批中”，不能重复提交。'),
    ('领导自己登记为什么立刻就生效？',
     '领导本身就是审批人，再让自己审自己没有意义，所以管理层登记直接落地。'),
    ('措施到期了怎么处理？',
     '在工作台或“我的案件”的“措施到期”清单里点“登记下一措施”，选新的措施和期限提交即可；'
     '承办人提交后提醒领导审批。'),
    ('嫌疑人删掉重建会不会丢记录？',
     '现在不需要删了。点行内“编辑”直接改，修改前后的完整快照都会留在操作日志里，可撤回。'),
    ('日期为什么不能填时分秒了？',
     '办案时间按天计算，时分秒只会造成误判（比如“今天到期”被算成逾期）。'
     '需要精确时间的地方只有操作日志，它仍然精确到秒。'),
    ('升级会不会影响已经录入的数据？',
     '不会。升级只新增可空的字段，不删不改；一次性数据整理有台账闸门，一个库只执行一次。'
     '升级前备份数据库是稳妥习惯，但不是必需品。'),
]
for i, (q, a) in enumerate(qa, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    cn(p.add_run(f'Q{i}  '), font='微软雅黑', size=10.5, bold=True, color=GOLD)
    cn(p.add_run(q), font='微软雅黑', size=10.5, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.9)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.5
    cn(p2.add_run('A  '), font='微软雅黑', size=10, bold=True, color=BLUE)
    cn(p2.add_run(a), size=10)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
cn(p.add_run('—— 以上内容对应升级补丁 20261025，如有疑问请联系系统维护人员 ——'),
   font='微软雅黑', size=9, color=GREY)

doc.save(OUT)
print('saved:', OUT)
print('size:', round(os.path.getsize(OUT) / 1024, 1), 'KB')
